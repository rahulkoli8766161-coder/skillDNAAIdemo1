"""
SkillDNA-AI Resume Parser Module
Extracts text content from PDF, DOCX, and TXT resume files.
"""

import os

try:
    from PyPDF2 import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document
except ImportError:
    Document = None


# Maximum file size: 5MB
MAX_FILE_SIZE = 5 * 1024 * 1024
ALLOWED_EXTENSIONS = {'pdf', 'docx', 'txt'}


def allowed_file(filename):
    """Check if the file extension is allowed."""
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def get_file_extension(filename):
    """Get the file extension in lowercase."""
    return filename.rsplit('.', 1)[1].lower() if '.' in filename else ''


def validate_file(filepath):
    """
    Validate the uploaded file.
    Returns (is_valid, message).
    """
    if not os.path.exists(filepath):
        return False, 'File not found.'

    # Check file size
    file_size = os.path.getsize(filepath)
    if file_size > MAX_FILE_SIZE:
        return False, f'File size ({file_size / (1024*1024):.1f}MB) exceeds the 5MB limit.'

    if file_size == 0:
        return False, 'File is empty.'

    # Check extension
    ext = get_file_extension(filepath)
    if ext not in ALLOWED_EXTENSIONS:
        return False, f'File format .{ext} is not supported. Please upload PDF, DOCX, or TXT.'

    return True, 'File is valid.'


def extract_text_from_pdf(filepath):
    """
    Extract text content from a PDF file.
    Returns extracted text string.
    """
    try:
        reader = PdfReader(filepath)
        text_parts = []

        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text.strip())

        full_text = '\n\n'.join(text_parts)

        if not full_text.strip():
            return '[PDF contained no extractable text. The file may be image-based.]'

        return full_text

    except Exception as e:
        return f'[Error extracting PDF text: {str(e)}]'


def extract_text_from_docx(filepath):
    """
    Extract text content from a DOCX file.
    Returns extracted text string.
    """
    try:
        doc = Document(filepath)
        text_parts = []

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text.strip())

        # Also extract from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = ' | '.join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_parts.append(row_text)

        full_text = '\n'.join(text_parts)

        if not full_text.strip():
            return '[DOCX contained no extractable text.]'

        return full_text

    except Exception as e:
        return f'[Error extracting DOCX text: {str(e)}]'


def extract_text_from_txt(filepath):
    """
    Extract text content from a TXT file.
    Returns extracted text string.
    """
    try:
        with open(filepath, 'r', encoding='utf-8', errors='replace') as f:
            text = f.read()

        if not text.strip():
            return '[TXT file is empty.]'

        return text

    except Exception as e:
        return f'[Error reading TXT file: {str(e)}]'


def parse_resume(filepath):
    """
    Auto-detect file format and extract text from resume.
    Returns (success, extracted_text_or_error_message).
    """
    # Validate file first
    is_valid, message = validate_file(filepath)
    if not is_valid:
        return False, message

    ext = get_file_extension(filepath)

    extractors = {
        'pdf': extract_text_from_pdf,
        'docx': extract_text_from_docx,
        'txt': extract_text_from_txt
    }

    extractor = extractors.get(ext)
    if not extractor:
        return False, f'Unsupported file format: .{ext}'

    try:
        text = extractor(filepath)
        return True, text
    except Exception as e:
        return False, f'Failed to parse resume: {str(e)}'
